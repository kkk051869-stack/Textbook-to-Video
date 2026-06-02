"""测试 DOCX 文档解析模块"""

import pytest
from docx import Document

from textbook2video.pipeline.docx_parser import (
    extract_full_text,
    extract_lesson_info,
    extract_section,
    list_lessons,
    list_sections,
)


@pytest.fixture
def docx_with_headings(tmp_path):
    """生成带标题结构的测试 DOCX 文件。"""
    path = tmp_path / "test_textbook.docx"
    doc = Document()

    doc.add_heading("第一章 人工智能概述", level=1)
    doc.add_paragraph("人工智能（AI）是计算机科学的一个重要分支。")
    doc.add_paragraph("本章将介绍人工智能的基本概念和发展历程。")
    doc.add_paragraph("AI 技术正在改变我们的生活。")

    doc.add_heading("第二章 机器学习基础", level=1)
    doc.add_paragraph("机器学习是实现人工智能的核心方法之一。")
    doc.add_paragraph("监督学习、无监督学习和强化学习是三种主要的学习范式。")

    doc.add_heading("2.1 监督学习", level=2)
    doc.add_paragraph("监督学习需要标注好的训练数据。")
    doc.add_paragraph("常见的算法包括线性回归和决策树。")

    doc.add_heading("2.2 无监督学习", level=2)
    doc.add_paragraph("无监督学习不需要标签。")
    doc.add_paragraph("聚类和降维是常见的无监督学习方法。")

    doc.add_heading("第三章 深度学习", level=1)
    doc.add_paragraph("深度学习使用多层神经网络来学习数据表示。")
    doc.add_paragraph("卷积神经网络在图像识别领域取得了突破性进展。")
    doc.add_paragraph("循环神经网络擅长处理序列数据。")

    doc.save(str(path))
    return str(path)


@pytest.fixture
def docx_no_headings(tmp_path):
    """生成无标题结构的测试 DOCX 文件。"""
    path = tmp_path / "plain_text.docx"
    doc = Document()
    doc.add_paragraph("这是一段普通文本。")
    doc.add_paragraph("这是第二段普通文本，没有任何标题。")
    doc.add_paragraph("第三段内容讲述了一些有趣的知识。")
    doc.save(str(path))
    return str(path)


@pytest.fixture
def docx_with_table(tmp_path):
    """生成带表格的测试 DOCX 文件。"""
    path = tmp_path / "with_table.docx"
    doc = Document()
    doc.add_heading("数据表格", level=1)
    doc.add_paragraph("以下是一些统计数据：")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "准确率"
    table.cell(1, 1).text = "95.2%"
    table.cell(2, 0).text = "召回率"
    table.cell(2, 1).text = "89.7%"
    doc.add_paragraph("表格展示完毕。")
    doc.save(str(path))
    return str(path)


class TestExtractFullText:
    """测试 extract_full_text"""

    def test_basic_extraction(self, docx_with_headings):
        text = extract_full_text(docx_with_headings)
        assert "人工智能" in text
        assert "机器学习" in text
        assert "深度学习" in text

    def test_includes_all_paragraphs(self, docx_with_headings):
        text = extract_full_text(docx_with_headings)
        assert "监督学习需要标注好的训练数据" in text
        assert "卷积神经网络" in text

    def test_table_extraction(self, docx_with_table):
        text = extract_full_text(docx_with_table)
        assert "准确率" in text
        assert "95.2%" in text

    def test_file_not_found(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            extract_full_text(str(tmp_path / "nonexistent.docx"))

    def test_invalid_extension(self, tmp_path):
        fake = tmp_path / "test.txt"
        fake.write_text("hello")
        with pytest.raises(ValueError, match="非 DOCX"):
            extract_full_text(str(fake))


class TestListSections:
    """测试 list_sections"""

    def test_section_count(self, docx_with_headings):
        """Heading 1 和 Heading 2 都作为分界（max_level=2）。"""
        sections = list_sections(docx_with_headings)
        # 第一章, 第二章, 2.1, 2.2, 第三章 = 5 sections
        assert len(sections) == 5

    def test_section_count_level1_only(self, docx_with_headings):
        """只按 Heading 1 拆分。"""
        sections = list_sections(docx_with_headings, max_level=1)
        # 第一章, 第二章, 第三章 = 3 sections
        assert len(sections) == 3

    def test_section_structure(self, docx_with_headings):
        sections = list_sections(docx_with_headings)
        for sec in sections:
            assert "section_number" in sec
            assert "title" in sec
            assert "para_count" in sec

    def test_section_titles(self, docx_with_headings):
        sections = list_sections(docx_with_headings)
        titles = [s["title"] for s in sections]
        assert "第一章 人工智能概述" in titles
        assert "第二章 机器学习基础" in titles
        assert "2.1 监督学习" in titles

    def test_no_headings_fallback(self, docx_no_headings):
        """无标题的文档整体作为一个章节。"""
        sections = list_sections(docx_no_headings)
        assert len(sections) == 1
        assert sections[0]["para_count"] == 3


class TestExtractSection:
    """测试 extract_section"""

    def test_extract_first_section(self, docx_with_headings):
        text = extract_section(docx_with_headings, 1)
        assert "人工智能" in text
        assert "基本概念和发展历程" in text

    def test_extract_subsection(self, docx_with_headings):
        """提取 2.1 监督学习子章节。"""
        sections = list_sections(docx_with_headings)
        # 找到 2.1 的编号
        for sec in sections:
            if "监督学习" in sec["title"]:
                text = extract_section(docx_with_headings, sec["section_number"])
                assert "标注好的训练数据" in text
                break

    def test_invalid_section_number(self, docx_with_headings):
        with pytest.raises(ValueError, match="不支持的章节号"):
            extract_section(docx_with_headings, 0)
        with pytest.raises(ValueError, match="不支持的章节号"):
            extract_section(docx_with_headings, 999)


class TestExtractLessonInfo:
    """测试 extract_lesson_info（对齐 PDF parser 接口）"""

    def test_info_structure(self, docx_with_headings):
        info = extract_lesson_info(docx_with_headings, 1)
        assert info["lesson_number"] == 1
        assert info["pages"] == (0, 0)
        assert "text" in info
        assert "title" in info

    def test_info_text_content(self, docx_with_headings):
        info = extract_lesson_info(docx_with_headings, 1)
        assert "人工智能" in info["text"]
        assert info["title"] == "第一章 人工智能概述"

    def test_info_invalid_section(self, docx_with_headings):
        with pytest.raises(ValueError, match="不支持的章节号"):
            extract_lesson_info(docx_with_headings, 0)
        with pytest.raises(ValueError, match="不支持的章节号"):
            extract_lesson_info(docx_with_headings, 999)


class TestListLessons:
    """测试 list_lessons（对齐 PDF parser 接口）"""

    def test_list_structure(self, docx_with_headings):
        lessons = list_lessons(docx_with_headings)
        for lesson in lessons:
            assert "lesson_number" in lesson
            assert "pages" in lesson
            assert "page_count" in lesson
            assert "title" in lesson
            assert "para_count" in lesson

    def test_list_count(self, docx_with_headings):
        """max_level=2 时，5 个章节（3 个 H1 + 2 个 H2）。"""
        lessons = list_lessons(docx_with_headings)
        assert len(lessons) == 5

    def test_pages_are_placeholder(self, docx_with_headings):
        """DOCX 没有页码概念，pages 固定为 (0,0)。"""
        lessons = list_lessons(docx_with_headings)
        for lesson in lessons:
            assert lesson["pages"] == (0, 0)
            assert lesson["page_count"] == 0

    def test_para_count(self, docx_with_headings):
        """各章节的段落数正确。"""
        lessons = list_lessons(docx_with_headings)
        # 第一章: 3 段正文
        assert lessons[0]["para_count"] == 3
        assert lessons[0]["title"] == "第一章 人工智能概述"

    def test_no_headings(self, docx_no_headings):
        lessons = list_lessons(docx_no_headings)
        assert len(lessons) == 1
        assert lessons[0]["para_count"] == 3
