#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
智能样式映射脚本 - 将 v6.docx 的段落映射到 moban.docx 的正确样式
"""
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
import re
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

moban_path = r'D:\text python\Textbook-to-Video-master\Textbook-to-Video-master\docs\moban.docx'
v6_path = r'D:\text python\Textbook-to-Video-master\Textbook-to-Video-master\docs\v6.docx'
output_path = r'D:\text python\Textbook-to-Video-master\Textbook-to-Video-master\docs\v6_formatted_final.docx'

moban = Document(moban_path)
v6 = Document(v6_path)

# 定义智能样式映射规则
def determine_style_for_paragraph(p, index):
    """
    根据段落内容智能判断应该使用 moban 中的哪个样式
    """
    text = p.text.strip() if p.text else ''
    
    # 特殊段落检测
    if index == 0:
        return 'Heading 2'  # 论文标题
    
    if '华东师范大学' in text and '数据科学与工程学院' in text:
        return '样式1'  # 作者单位
    
    if 'SHUANG He' in text or 'East China Normal University' in text:
        return 'Heading 5' if 'SHUANG He' in text else '样式3'
    
    if '摘 要' in text or 'Abstract:' in text:
        return '样式2' if '摘 要' in text else '样式4'
    
    if '关键词' in text or 'Key words:' in text:
        return '样式2' if '关键词' in text else 'Heading 6'
    
    # 标题检测
    if re.match(r'^\d+\s+.+$', text):  # 一级标题
        return 'Heading 7'
    
    if re.match(r'^\d+\.\d+\s+.+$', text):  # 二级标题
        return 'Heading 8'
    
    if re.match(r'^\d+\.\d+\.\d+\s+.+$', text):  # 三级标题
        return '标题9'
    
    if '参考文献' in text:
        return '参考文献'
    
    # 图表标题
    if text.startswith('图') or text.startswith('Fig'):
        return '样式1'  # 图题使用样式1
    
    if text.startswith('表') or text.startswith('Table'):
        return '样式1'  # 表题使用样式1
    
    # 默认正文
    return 'Normal'

# 1. 首先确保 v6 中有所有需要的样式
print("检查并创建样式...")
v6_all_styles = {s.name for s in v6.styles}
moban_styles = {s.name: s for s in moban.styles}

# 需要创建的新样式
styles_to_create = []
for sn in ['Heading 2', 'Heading 6', 'Heading 7', 'Heading 8', '参考文献', 'Normal (Web)', '样式2', '样式4']:
    if sn not in v6_all_styles:
        styles_to_create.append(sn)

for sn in styles_to_create:
    ms = moban_styles[sn]
    ns = v6.styles.add_style(sn, WD_STYLE_TYPE.PARAGRAPH)
    
    if ms.font.name:
        ns.font.name = ms.font.name
    if ms.font.size:
        ns.font.size = ms.font.size
    if ms.font.bold is not None:
        ns.font.bold = ms.font.bold
    if ms.font.italic is not None:
        ns.font.italic = ms.font.italic
    
    if ms.paragraph_format.space_before is not None:
        ns.paragraph_format.space_before = ms.paragraph_format.space_before
    if ms.paragraph_format.space_after is not None:
        ns.paragraph_format.space_after = ms.paragraph_format.space_after
    if ms.paragraph_format.line_spacing is not None:
        ns.paragraph_format.line_spacing = ms.paragraph_format.line_spacing
    if ms.paragraph_format.alignment is not None:
        ns.paragraph_format.alignment = ms.paragraph_format.alignment
    if ms.paragraph_format.first_line_indent is not None:
        ns.paragraph_format.first_line_indent = ms.paragraph_format.first_line_indent
    
    print(f"已创建样式: {sn}")

# 2. 复制共同样式的格式
print("\n复制共同样式格式...")
common = set(p.style.name for p in moban.paragraphs if p.style) & set(s.name for s in v6.styles)
for sn in common:
    ms = moban_styles[sn]
    ds = v6.styles[sn]
    
    if ms.font.name:
        ds.font.name = ms.font.name
    if ms.font.size:
        ds.font.size = ms.font.size
    if ms.font.bold is not None:
        ds.font.bold = ms.font.bold
    if ms.font.italic is not None:
        ds.font.italic = ms.font.italic
    
    if ms.paragraph_format.space_before is not None:
        ds.paragraph_format.space_before = ms.paragraph_format.space_before
    if ms.paragraph_format.space_after is not None:
        ds.paragraph_format.space_after = ms.paragraph_format.space_after
    if ms.paragraph_format.line_spacing is not None:
        ds.paragraph_format.line_spacing = ms.paragraph_format.line_spacing
    if ms.paragraph_format.alignment is not None:
        ds.paragraph_format.alignment = ms.paragraph_format.alignment
    if ms.paragraph_format.first_line_indent is not None:
        ds.paragraph_format.first_line_indent = ms.paragraph_format.first_line_indent

# 3. 重新映射段落样式
print("\n重新映射段落样式...")
style_map = {}
for i, p in enumerate(v6.paragraphs):
    target_style = determine_style_for_paragraph(p, i)
    
    # 检查目标样式是否存在
    if target_style not in [s.name for s in v6.styles]:
        print(f"警告: 目标样式 {target_style} 不存在")
        continue
    
    # 应用新样式
    p.style = v6.styles[target_style]
    style_map[i] = target_style
    
    text = p.text[:60] if p.text else '(empty)'
    print(f"[{i:2d}] {p.style.name:20s} {text}")

# 4. 复制页面设置
print("\n复制页面设置...")
for section_idx in range(len(moban.sections)):
    ms = moban.sections[section_idx]
    if section_idx < len(v6.sections):
        ds = v6.sections[section_idx]
    else:
        ds = v6.sections.add_section()
    
    ds.top_margin = ms.top_margin
    ds.bottom_margin = ms.bottom_margin
    ds.left_margin = ms.left_margin
    ds.right_margin = ms.right_margin
    ds.page_width = ms.page_width
    ds.page_height = ms.page_height

# 5. 保存
print(f"\n保存文件...")
v6.save(output_path)
print(f"转换完成！已保存到: {output_path}")

# 6. 验证
result = Document(output_path)
print("\n=== 验证结果 ===")
print(f"段落数: {len(result.paragraphs)}")

print("\n验证样式属性:")
for sn in ['Heading 2', 'Heading 5', 'Heading 7', 'Heading 8', 'Normal', '样式1', '样式3', '参考文献', 'Normal (Web)', '样式2', '样式4']:
    try:
        s = result.styles[sn]
        print(f"{sn}: 字号={Pt(s.font.size).pt if s.font.size else '无'}, 行距={s.paragraph_format.line_spacing}, 对齐={s.paragraph_format.alignment}")
    except:
        print(f"{sn}: 未找到")
