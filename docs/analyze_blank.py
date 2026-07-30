#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Detailed analysis of all paragraphs in moban.docx
"""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

doc = Document(r'D:\text python\Textbook-to-Video-master\Textbook-to-Video-master\docs\moban.docx')

print(f"总段落数: {len(doc.paragraphs)}")
print(f"总节数: {len(doc.sections)}")

print("\n=== 所有段落详细分析 ===")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    is_empty = (text == '' or text.isspace())
    sn = p.style.name if p.style else 'None'
    
    # 检查分页符
    has_page_break = False
    has_section_break = False
    pPr = p._element.find(qn('w:pPr'))
    if pPr is not None:
        br = pPr.find(qn('w:br'))
        if br is not None:
            type_attr = br.get(qn('w:type'))
            if type_attr == 'page':
                has_page_break = True
    
    # 检查段落后是否有 sectPr (分节符)
    next_elem = p._element.getnext()
    if next_elem is not None:
        tag = next_elem.tag.split('}')[-1] if '}' in next_elem.tag else next_elem.tag
        if tag == 'sectPr':
            has_section_break = True
    
    display = '(空)' if is_empty else text[:70]
    flags = []
    if has_page_break:
        flags.append('分页符')
    if has_section_break:
        flags.append('分节符')
    flag_str = f' [{", ".join(flags)}]' if flags else ''
    
    print(f"[{i:2d}] {sn:18s} {display}{flag_str}")
