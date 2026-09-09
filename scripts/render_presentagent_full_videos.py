from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
from pathlib import Path

from pptx import Presentation


def render_pdf_frames(pdf_path: Path, frames_dir: Path, env: dict[str, str]) -> list[Path]:
    pdftoppm = shutil.which("pdftoppm", path=env.get("PATH"))
    if pdftoppm:
        subprocess.run(
            [pdftoppm, "-jpeg", "-r", "160", str(pdf_path), str(frames_dir / "frame")],
            check=True,
            env=env,
        )
        return sorted(frames_dir.glob("frame-*.jpg"))

    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf_path))
    zoom = 160 / 72
    frames: list[Path] = []
    for idx, page in enumerate(doc, start=1):
        frame = frames_dir / f"frame-{idx}.jpg"
        image = page.render(scale=zoom).to_pil()
        image.save(frame, format="JPEG", quality=92)
        frames.append(frame)
    return frames


def generate_megatts3_audio(segments: list[str], output_dir: Path) -> None:
    python = os.getenv("T2V_MEGATTS3_PYTHON", "/ai/data/tools/envs/cosyvoice3/bin/python")
    megatts3_root = Path(os.getenv("T2V_MEGATTS3_ROOT", "/ai/data/repos/PresentAgent/presentagent/MegaTTS3"))
    runner = Path("/ai/data/repos/Textbook-to-Video/src/textbook2video/pipeline/megatts3_runner.py")
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join(
        value for value in (str(megatts3_root), env.get("PYTHONPATH", "")) if value
    )
    env["PATH"] = os.pathsep.join(
        value
        for value in (str(Path(python).parent), "/ai/data/tools/bin", "/usr/bin", env.get("PATH", ""))
        if value
    )
    request = {
        "segments": segments,
        "output_dir": str(output_dir),
        "megatts3_root": str(megatts3_root),
        "prompt_wav": str(megatts3_root / "assets" / "Chinese_prompt.wav"),
        "prompt_latent": str(megatts3_root / "assets" / "Chinese_prompt.npy"),
        "time_step": int(os.getenv("T2V_MEGATTS3_TIME_STEP", "16")),
        "p_w": float(os.getenv("T2V_MEGATTS3_P_W", "1.6")),
        "t_w": float(os.getenv("T2V_MEGATTS3_T_W", "2.5")),
        "device": os.getenv("T2V_MEGATTS3_DEVICE", "cuda"),
    }
    subprocess.run(
        [python, str(runner)],
        input=json.dumps(request, ensure_ascii=False),
        text=True,
        encoding="utf-8",
        env=env,
        check=True,
    )


def slide_notes_or_text(pptx_path: Path) -> list[str]:
    prs = Presentation(str(pptx_path))
    notes: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        text = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not text:
            slide_texts = []
            for shape in slide.shapes:
                shape_text = getattr(shape, "text", "")
                if shape_text and shape_text.strip():
                    slide_texts.append(shape_text.strip())
            text = "。".join(slide_texts) or f"这是第 {idx} 页。"
        notes.append(text)
    return notes


def render_one(run_dir: Path, force: bool) -> None:
    pptx_path = run_dir / "generated_full_presentagent.pptx"
    if not pptx_path.exists():
        raise FileNotFoundError(pptx_path)
    task_dir = run_dir / "presentagent_video_task"
    output = task_dir / "output.mp4"
    if output.exists() and output.stat().st_size > 0 and not force:
        print(f"SKIP {run_dir.name}: {output}", flush=True)
        return

    if task_dir.exists() and force:
        shutil.rmtree(task_dir)
    task_dir.mkdir(parents=True, exist_ok=True)
    task_pptx = task_dir / "source.pptx"
    shutil.copy2(pptx_path, task_pptx)

    env = os.environ.copy()
    env["PATH"] = "/usr/bin:/ai/data/tools/libreoffice/program:/ai/data/tools/bin:" + env.get("PATH", "")
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", str(task_pptx), "--outdir", str(task_dir)],
        check=True,
        env=env,
    )
    pdf_path = task_dir / "source.pdf"
    frames_dir = task_dir / "frames"
    audio_dir = task_dir / "audio"
    segments_dir = task_dir / "segments"
    for directory in (frames_dir, audio_dir, segments_dir):
        directory.mkdir(parents=True, exist_ok=True)

    frames = render_pdf_frames(pdf_path, frames_dir, env)
    notes = slide_notes_or_text(task_pptx)
    if len(frames) != len(notes):
        raise RuntimeError(f"{run_dir.name}: frame count {len(frames)} != notes count {len(notes)}")

    (task_dir / "notes.json").write_text(json.dumps(notes, ensure_ascii=False, indent=2), encoding="utf-8")
    generate_megatts3_audio(notes, audio_dir)

    segment_paths = []
    for idx, frame in enumerate(frames, start=1):
        audio = audio_dir / f"s{idx}.wav"
        if not audio.exists() or audio.stat().st_size == 0:
            raise RuntimeError(f"missing audio: {audio}")
        segment = segments_dir / f"segment_{idx:03d}.mp4"
        subprocess.run(
            [
                "ffmpeg",
                "-y",
                "-loop",
                "1",
                "-i",
                str(frame),
                "-i",
                str(audio),
                "-vf",
                "scale=1920:1080:force_original_aspect_ratio=decrease,pad=1920:1080:(ow-iw)/2:(oh-ih)/2,format=yuv420p",
                "-c:v",
                "libx264",
                "-tune",
                "stillimage",
                "-c:a",
                "aac",
                "-b:a",
                "192k",
                "-shortest",
                str(segment),
            ],
            check=True,
            env=env,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        segment_paths.append(segment)

    list_file = task_dir / "segments.txt"
    list_file.write_text("\n".join(f"file '{path.as_posix()}'" for path in segment_paths) + "\n", encoding="utf-8")
    subprocess.run(
        ["ffmpeg", "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", str(output)],
        check=True,
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    print(f"DONE {run_dir.name}: {output}", flush=True)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("roots", nargs="+")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    failures = 0
    for root_arg in args.roots:
        root = Path(root_arg)
        for run_dir in sorted(p for p in root.iterdir() if p.is_dir()):
            try:
                render_one(run_dir, args.force)
            except Exception as exc:
                failures += 1
                print(f"FAIL {run_dir}: {exc}", flush=True)
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
