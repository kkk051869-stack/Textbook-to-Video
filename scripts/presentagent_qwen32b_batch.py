from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import shutil
import subprocess
import sys
import wave
from dataclasses import dataclass
from pathlib import Path

from openai import OpenAI
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass
class Item:
    item_id: str
    kind: str
    title: str
    source_path: Path
    output_dir: Path
    slides: int


def safe_id(text: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_")
    return value[:80] or "item"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def extract_docx_text(path: Path, max_chars: int) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                chunks.append(line)
    return "\n".join(chunks)[:max_chars]


def extract_source_json_text(path: Path, max_chars: int) -> tuple[str, str]:
    data = read_json(path)
    title = data.get("title") or data.get("section_title") or data.get("lesson_id") or path.parent.name
    chunks: list[str] = []
    for paragraph in data.get("paragraphs", []):
        text = str(paragraph.get("text", "")).strip()
        if text:
            chunks.append(text)
    for image in data.get("images", []):
        caption = str(image.get("caption", "")).strip()
        if caption:
            chunks.append(f"Figure: {caption}")
    return title, "\n".join(chunks)[:max_chars]


def extract_pdf_text(path: Path, max_chars: int) -> str:
    result = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", "-layout", str(path), "-"],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return result.stdout.decode("utf-8", errors="replace")[:max_chars]


def extract_json_array(text: str):
    try:
        import json_repair
    except Exception:
        json_repair = None

    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S).strip()
    fenced = re.search(r"```(?:json)?\s*(\[.*?\])\s*```", text, flags=re.S)
    if fenced:
        text = fenced.group(1)
    else:
        start = text.find("[")
        end = text.rfind("]")
        if start != -1 and end != -1 and end > start:
            text = text[start : end + 1]
    if json_repair is not None:
        return json_repair.loads(text)
    return json.loads(text)


def make_slide_plan(title: str, text: str, slides: int, api_base: str, model: str) -> list[dict]:
    client = OpenAI(base_url=api_base, api_key="EMPTY", timeout=900)
    prompt = f"""
Create a Chinese teaching-video slide plan from the textbook excerpt below.

Return ONLY a JSON array. Do not use Markdown.
The array must contain exactly {slides} objects.
Each object must have:
- title: concise Chinese slide title
- bullets: 3 to 5 short Chinese bullet points
- speaker_notes: Chinese narration text, 80 to 140 Chinese characters

Rules:
- Use only information supported by the textbook excerpt.
- Do not add outside facts.
- Prefer clear teaching structure over decorative wording.
- Keep the narration suitable for text-to-speech.

Lesson title:
{title}

Textbook excerpt:
{text}
"""
    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": "You are a careful instructional slide designer. Output valid JSON only.",
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0.2,
        max_tokens=2400,
        extra_body={"chat_template_kwargs": {"enable_thinking": False}},
    )
    content = response.choices[0].message.content or ""
    plan = extract_json_array(content)
    if not isinstance(plan, list):
        raise ValueError("Model output is not a JSON array")
    return normalize_plan(plan, slides, title)


def normalize_plan(plan: list, slides: int, fallback_title: str) -> list[dict]:
    normalized: list[dict] = []
    for idx, item in enumerate(plan[:slides], start=1):
        if not isinstance(item, dict):
            item = {}
        title = str(item.get("title") or f"{fallback_title} {idx}").strip()
        bullets = item.get("bullets") or []
        if isinstance(bullets, str):
            bullets = [line.strip("-* ") for line in bullets.splitlines() if line.strip()]
        bullets = [str(b).strip() for b in bullets if str(b).strip()][:5]
        while len(bullets) < 3:
            bullets.append(title)
        notes = str(item.get("speaker_notes") or "，".join(bullets)).strip()
        normalized.append({"title": title, "bullets": bullets, "speaker_notes": notes})
    while len(normalized) < slides:
        normalized.append(
            {
                "title": f"{fallback_title} {len(normalized) + 1}",
                "bullets": [fallback_title, "关键概念", "学习要点"],
                "speaker_notes": f"这一页围绕{fallback_title}的关键内容展开，帮助学习者建立基本理解。",
            }
        )
    return normalized


def set_run_font(run, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        try:
            run.font.color.rgb = color
        except Exception:
            pass


def add_textbox(slide, left, top, width, height, text: str, font_size: int, bold=False, color=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_run_font(run, font_size, bold, color)
    return box


def build_pptx(plan: list[dict], pptx_path: Path, lesson_title: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_color = RGBColor(20, 33, 61)
    text_color = RGBColor(31, 41, 55)

    for idx, item in enumerate(plan, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
        try:
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor(29, 78, 216)
            band.line.fill.background()
        except Exception:
            pass

        add_textbox(slide, Inches(0.65), Inches(0.45), Inches(12.0), Inches(0.85), item["title"], 30, True, title_color)

        body_box = slide.shapes.add_textbox(Inches(0.95), Inches(1.6), Inches(11.45), Inches(4.95))
        frame = body_box.text_frame
        frame.word_wrap = True
        frame.clear()
        for bullet_idx, bullet in enumerate(item["bullets"][:5]):
            p = frame.paragraphs[0] if bullet_idx == 0 else frame.add_paragraph()
            p.text = str(bullet)
            p.level = 0
            p.space_after = Pt(10)
            for run in p.runs:
                set_run_font(run, 22, False, text_color)

        footer = f"{lesson_title}  |  {idx}/{len(plan)}"
        add_textbox(slide, Inches(0.7), Inches(6.9), Inches(12), Inches(0.3), footer, 10, False, RGBColor(100, 116, 139))
        slide.notes_slide.notes_text_frame.text = item["speaker_notes"]

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(pptx_path)


def item_text(item: Item, max_chars: int) -> str:
    if item.kind == "docx":
        return extract_docx_text(item.source_path, max_chars)
    if item.kind == "lesson":
        _title, text = extract_source_json_text(item.source_path, max_chars)
        return text
    if item.source_path.suffix.lower() == ".pdf":
        return extract_pdf_text(item.source_path, max_chars)
    raise ValueError(f"Unsupported item: {item}")


def generate_content(item: Item, args: argparse.Namespace) -> None:
    item.output_dir.mkdir(parents=True, exist_ok=True)
    plan_path = item.output_dir / "qwen32b_slide_plan.json"
    pptx_path = item.output_dir / "generated.pptx"
    excerpt_path = item.output_dir / "source_excerpt.txt"
    status_path = item.output_dir / "run_status.json"

    if plan_path.exists() and pptx_path.exists() and not args.force:
        print(f"SKIP content {item.item_id}: existing PPT/plan", flush=True)
        return

    text = item_text(item, args.max_chars)
    excerpt_path.write_text(text, encoding="utf-8")
    if plan_path.exists() and not args.force:
        plan = read_json(plan_path)
        print(f"REUSE plan {item.item_id}: {plan_path}", flush=True)
    else:
        plan = make_slide_plan(item.title, text, item.slides, args.api_base, args.model)
        write_json(plan_path, plan)
    build_pptx(plan, pptx_path, item.title)
    write_json(
        status_path,
        {
            "item_id": item.item_id,
            "kind": item.kind,
            "status": "content_completed",
            "title": item.title,
            "slides": item.slides,
            "source_path": str(item.source_path),
            "slide_plan": str(plan_path),
            "generated_pptx": str(pptx_path),
            "model": args.model,
            "api_base": args.api_base,
        },
    )
    print(f"DONE content {item.item_id}: {pptx_path}", flush=True)


async def render_video_with_presentagent(item: Item, args: argparse.Namespace) -> None:
    pptx_path = item.output_dir / "generated.pptx"
    if not pptx_path.exists():
        raise FileNotFoundError(pptx_path)

    task_dir = item.output_dir / "presentagent_video_task"
    task_dir.mkdir(parents=True, exist_ok=True)
    output = task_dir / "output.mp4"
    if output.exists() and output.stat().st_size > 0 and not args.force:
        print(f"SKIP video {item.item_id}: existing {output}", flush=True)
        return
    task_pptx = task_dir / "source.pptx"
    shutil.copy2(pptx_path, task_pptx)

    env = os.environ.copy()
    env["PATH"] = "/usr/bin:/ai/data/tools/libreoffice/program:/ai/data/tools/bin:" + env.get("PATH", "")
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            str(task_pptx),
            "--outdir",
            str(task_dir),
        ],
        check=True,
        env=env,
    )
    pdf_path = task_dir / "source.pdf"
    frames_dir = task_dir / "frames"
    audio_dir = task_dir / "audio"
    segments_dir = task_dir / "segments"
    for directory in (frames_dir, audio_dir, segments_dir):
        if directory.exists():
            shutil.rmtree(directory)
        directory.mkdir(parents=True, exist_ok=True)

    subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "160", str(pdf_path), str(frames_dir / "frame")],
        check=True,
        env=env,
    )
    frames = sorted(frames_dir.glob("frame-*.jpg"))

    prs = Presentation(str(task_pptx))
    notes: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        text = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        notes.append(text or f"这是第 {idx} 页。")
    if len(frames) != len(notes):
        raise RuntimeError(f"Frame count {len(frames)} does not match slide count {len(notes)}")

    generate_megatts3_audio(notes, audio_dir)
    audio_files = [audio_dir / f"s{idx}.wav" for idx in range(1, len(notes) + 1)]
    missing_audio = [str(path) for path in audio_files if not path.exists() or path.stat().st_size == 0]
    if missing_audio:
        raise RuntimeError(f"MegaTTS3 did not produce audio files: {missing_audio}")

    segment_paths: list[Path] = []
    for idx, (frame, audio) in enumerate(zip(frames, audio_files), start=1):
        segment = segments_dir / f"segment_{idx:03d}.mp4"
        subprocess.run(
            [
                "ffmpeg",
                "-y",
                "-loop",
                "1",
                "-i",
                str(frame),
                "-i",
                str(audio),
                "-vf",
                "scale=1920:1080:force_original_aspect_ratio=decrease,pad=1920:1080:(ow-iw)/2:(oh-ih)/2,format=yuv420p",
                "-c:v",
                "libx264",
                "-tune",
                "stillimage",
                "-c:a",
                "aac",
                "-b:a",
                "192k",
                "-shortest",
                str(segment),
            ],
            check=True,
            env=env,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        segment_paths.append(segment)

    list_file = task_dir / "segments.txt"
    list_file.write_text(
        "\n".join(f"file '{path.as_posix()}'" for path in segment_paths) + "\n",
        encoding="utf-8",
    )
    subprocess.run(
        ["ffmpeg", "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", str(output)],
        check=True,
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    print(f"DONE video {item.item_id}: {output}", flush=True)


def generate_megatts3_audio(segments: list[str], output_dir: Path) -> None:
    python = os.getenv("T2V_MEGATTS3_PYTHON", "/ai/data/tools/envs/cosyvoice3/bin/python")
    megatts3_root = Path(os.getenv("T2V_MEGATTS3_ROOT", "/ai/data/repos/PresentAgent/presentagent/MegaTTS3"))
    runner = Path("/ai/data/repos/Textbook-to-Video/src/textbook2video/pipeline/megatts3_runner.py")
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join(
        value for value in (str(megatts3_root), env.get("PYTHONPATH", "")) if value
    )
    env["PATH"] = os.pathsep.join(
        value for value in (str(Path(python).parent), "/ai/data/tools/bin", "/usr/bin", env.get("PATH", "")) if value
    )
    request = {
        "segments": segments,
        "output_dir": str(output_dir),
        "megatts3_root": str(megatts3_root),
        "prompt_wav": str(megatts3_root / "assets" / "Chinese_prompt.wav"),
        "prompt_latent": str(megatts3_root / "assets" / "Chinese_prompt.npy"),
        "time_step": int(os.getenv("T2V_MEGATTS3_TIME_STEP", "16")),
        "p_w": float(os.getenv("T2V_MEGATTS3_P_W", "1.6")),
        "t_w": float(os.getenv("T2V_MEGATTS3_T_W", "2.5")),
        "device": os.getenv("T2V_MEGATTS3_DEVICE", "cuda"),
    }
    subprocess.run(
        [python, str(runner)],
        input=json.dumps(request, ensure_ascii=False),
        text=True,
        encoding="utf-8",
        env=env,
        check=True,
    )


def collect_items(args: argparse.Namespace) -> list[Item]:
    items: list[Item] = []
    output_root = Path(args.output_root)

    if args.include in {"all", "docx"}:
        docx_root = Path(args.docx_root)
        for index, docx_path in enumerate(sorted(docx_root.rglob("*.docx")), start=1):
            title = docx_path.stem
            item_id = f"cs_chapter_{index:02d}_{safe_id(docx_path.stem)}"
            items.append(
                Item(
                    item_id=item_id,
                    kind="docx",
                    title=title,
                    source_path=docx_path,
                    output_dir=output_root / "computer_science_lecture" / item_id,
                    slides=args.docx_slides,
                )
            )

    if args.include in {"all", "lessons"}:
        experiment_root = Path(args.experiment_root)
        source_root = experiment_root / "sources" / "frozen"
        lesson_ids = args.lessons or sorted(p.name for p in source_root.iterdir() if (p / "source.json").exists())
        for lesson_id in lesson_ids:
            source_json = source_root / lesson_id / "source.json"
            title, _text = extract_source_json_text(source_json, args.max_chars)
            items.append(
                Item(
                    item_id=lesson_id,
                    kind="lesson",
                    title=title,
                    source_path=source_json,
                    output_dir=experiment_root / "runs" / "presentagent" / lesson_id,
                    slides=args.lesson_slides,
                )
            )
    return items


async def main_async() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=["content", "video", "all"], default="content")
    parser.add_argument("--include", choices=["all", "docx", "lessons"], default="all")
    parser.add_argument("--docx-root", default="/ai/data/textbook-to-video/inputs/presentagent-video-demo/extracted_utf8")
    parser.add_argument("--experiment-root", default="/ai/data/textbook-to-video/experiments/presentagent-comparison-v1")
    parser.add_argument("--output-root", default="/ai/data/textbook-to-video/outputs/presentagent-qwen32b-batch")
    parser.add_argument("--presentagent-repo", default="/ai/data/repos/PresentAgent")
    parser.add_argument("--api-base", default="http://127.0.0.1:8000/v1")
    parser.add_argument("--model", default="qwen3-32b-awq")
    parser.add_argument("--docx-slides", type=int, default=6)
    parser.add_argument("--lesson-slides", type=int, default=7)
    parser.add_argument("--max-chars", type=int, default=8000)
    parser.add_argument("--lessons", nargs="*", default=None)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    items = collect_items(args)
    print(f"Items: {len(items)}", flush=True)
    failures = 0
    for item in items:
        try:
            if args.stage in {"content", "all"}:
                generate_content(item, args)
            if args.stage in {"video", "all"}:
                await render_video_with_presentagent(item, args)
        except Exception as exc:
            failures += 1
            item.output_dir.mkdir(parents=True, exist_ok=True)
            write_json(
                item.output_dir / "run_status.json",
                {
                    "item_id": item.item_id,
                    "kind": item.kind,
                    "status": "failed",
                    "error": str(exc),
                    "source_path": str(item.source_path),
                },
            )
            print(f"FAIL {item.item_id}: {exc}", flush=True)
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main_async()))
