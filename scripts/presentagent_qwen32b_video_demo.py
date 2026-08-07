from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import shutil
import sys
import wave
from pathlib import Path

from docx import Document as DocxDocument
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt


def collect_docx_text(path: Path, max_chars: int) -> str:
    doc = DocxDocument(path)
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                chunks.append(line)
    text = "\n".join(chunks)
    return text[:max_chars]


def choose_docx(input_root: Path) -> Path:
    docs = sorted(input_root.rglob("*.docx"), key=lambda p: p.name)
    if not docs:
        raise FileNotFoundError(f"No .docx files found under {input_root}")
    for path in docs:
        if "\u7b2c\u4e00\u7ae0" in path.name:
            return path
    return docs[0]


def extract_json_array(text: str) -> list[dict]:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)
    fenced = re.search(r"```(?:json)?\s*(\[.*?\])\s*```", text, flags=re.S)
    if fenced:
        text = fenced.group(1)
    else:
        start = text.find("[")
        end = text.rfind("]")
        if start == -1 or end == -1 or end <= start:
            raise ValueError(f"Could not find JSON array in model response:\n{text[:1000]}")
        text = text[start : end + 1]
    return json.loads(text)


def make_slide_plan(text: str, slides: int, api_base: str, model: str) -> list[dict]:
    client = OpenAI(base_url=api_base, api_key="EMPTY", timeout=600)
    prompt = f"""
你正在把一章中文计算机科学讲义改写成教学视频用的幻灯片。
请根据教材正文生成 {slides} 页幻灯片。要求：
1. 只输出 JSON 数组，不要 Markdown，不要解释。
2. 每个元素包含 title, bullets, speaker_notes。
3. bullets 是 3 到 5 条短要点。
4. speaker_notes 是适合配音朗读的中文讲稿，80 到 140 字。
5. 内容必须来自教材，不要添加教材外的历史、事实或例子。

教材正文：
{text}
"""
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": "你是严谨的中文教材课件设计助手。输出必须是可解析 JSON。"},
            {"role": "user", "content": prompt},
        ],
        temperature=0.2,
        max_tokens=2200,
        extra_body={"chat_template_kwargs": {"enable_thinking": False}},
    )
    content = response.choices[0].message.content or ""
    return extract_json_array(content)


def add_textbox(slide, left, top, width, height, text: str, font_size: int, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return box


def build_pptx(plan: list[dict], pptx_path: Path, lesson_title: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for idx, item in enumerate(plan, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = str(item.get("title") or f"{lesson_title} - {idx}")
        bullets = item.get("bullets") or []
        notes = str(item.get("speaker_notes") or title)

        add_textbox(slide, Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.75), title, 30, True)
        body = "\n".join(f"\u2022 {str(b)}" for b in bullets[:5])
        add_textbox(slide, Inches(1.0), Inches(1.45), Inches(11.4), Inches(4.9), body, 22)
        slide.notes_slide.notes_text_frame.text = notes

    prs.save(pptx_path)


async def render_video_with_presentagent(pptx_path: Path, output_dir: Path, pa_repo: Path) -> Path:
    sys.path.insert(0, str(pa_repo))
    os.chdir(pa_repo)
    os.environ.setdefault("OPENAI_API_KEY", "EMPTY")
    os.environ.setdefault("API_BASE", "http://127.0.0.1:8000/v1")
    os.environ.setdefault("LANGUAGE_MODEL", "qwen3-32b-awq")
    os.environ.setdefault("VISION_MODEL", "qwen3-32b-awq")
    os.environ.setdefault("TEXT_MODEL", "qwen3-32b-awq")

    wrapper_dir = output_dir / "bin"
    wrapper_dir.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or "/ai/data/tools/libreoffice/program/soffice"
    libreoffice = wrapper_dir / "libreoffice"
    if not libreoffice.exists():
        libreoffice.symlink_to(soffice)
    os.environ["PATH"] = f"{wrapper_dir}:/ai/data/tools/libreoffice/program:/ai/data/tools/bin:" + os.environ.get("PATH", "")

    from presentagent import backend

    async def edge_tts_audio(text: str, output_path: str) -> None:
        try:
            import edge_tts

            mp3_path = str(Path(output_path).with_suffix(".mp3"))
            communicate = edge_tts.Communicate(text, voice="zh-CN-XiaoxiaoNeural")
            await communicate.save(mp3_path)
            await backend.run_cmd([
                "ffmpeg",
                "-y",
                "-i",
                mp3_path,
                "-ar",
                "22050",
                "-ac",
                "1",
                output_path,
            ])
        except Exception as exc:
            print(f"edge-tts failed, writing silence: {exc}")
            import numpy as np

            sample_rate = 22050
            duration = 4.0
            samples = np.zeros(int(sample_rate * duration), dtype=np.int16)
            with wave.open(output_path, "w") as wav_file:
                wav_file.setnchannels(1)
                wav_file.setsampwidth(2)
                wav_file.setframerate(sample_rate)
                wav_file.writeframes(samples.tobytes())

    backend.generate_tts_audio = edge_tts_audio

    task_dir = output_dir / "presentagent_video_task"
    task_dir.mkdir(parents=True, exist_ok=True)
    task_pptx = task_dir / "source.pptx"
    shutil.copy2(pptx_path, task_pptx)

    task_id = "qwen32b-demo"
    backend.ppt_video_progress_store[task_id] = {
        "task_dir": str(task_dir),
        "ppt_path": str(task_pptx),
        "status": "processing",
        "current_step": 0,
        "total_steps": 3,
        "current_slide": 0,
        "total_slides": 0,
        "progress_percentage": 0.0,
    }
    await backend.process_ppt_to_video(task_id)
    progress = backend.ppt_video_progress_store[task_id]
    if progress.get("status") != "completed":
        raise RuntimeError(f"PresentAgent video generation failed: {progress}")
    return task_dir / "output.mp4"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--presentagent-repo", type=Path, default=Path("/ai/data/repos/PresentAgent"))
    parser.add_argument("--api-base", default="http://127.0.0.1:8000/v1")
    parser.add_argument("--model", default="qwen3-32b-awq")
    parser.add_argument("--slides", type=int, default=6)
    parser.add_argument("--max-chars", type=int, default=8000)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = choose_docx(args.input_root)
    print(f"Using DOCX: {docx_path}")
    text = collect_docx_text(docx_path, args.max_chars)
    (args.output_dir / "source_excerpt.txt").write_text(text, encoding="utf-8")

    plan_path = args.output_dir / "qwen32b_slide_plan.json"
    if plan_path.exists():
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        print(f"Reusing slide plan: {plan_path}")
    else:
        plan = make_slide_plan(text, args.slides, args.api_base, args.model)
        plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Wrote slide plan: {plan_path}")

    pptx_path = args.output_dir / "qwen32b_presentagent_demo.pptx"
    if pptx_path.exists():
        print(f"Reusing PPTX: {pptx_path}")
    else:
        build_pptx(plan, pptx_path, docx_path.stem)
        print(f"Wrote PPTX: {pptx_path}")

    video_path = asyncio.run(render_video_with_presentagent(pptx_path, args.output_dir, args.presentagent_repo))
    print(f"Wrote video: {video_path}")


if __name__ == "__main__":
    main()
